# -*- coding: utf-8 -*-
"""
BOT FOTOS ITO + REPORTES FLASH EL6
Versión: v3 Operacional
Funciones:
- Registro de fotos con metadata profesional
- Reportes Flash
- Excel real .xlsx
- OneDrive vía Microsoft Graph
- Menú de botones
- Dashboard Telegram
- Buscar registros
- Estadísticas
- Últimos Flash
- Usuarios activos
"""

import os
import re
import logging
import threading
import asyncio
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
from http.server import BaseHTTPRequestHandler, HTTPServer
from collections import Counter

import requests
import msal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from telegram import Update, InlineKeyboardMarkup, InlineKeyboardButton
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    ConversationHandler,
    CallbackQueryHandler,
    ContextTypes,
    filters,
)

BOT_TOKEN = os.getenv("BOT_TOKEN", "")
if not BOT_TOKEN:
    raise RuntimeError("Define BOT_TOKEN en Render.")

CHILE_TZ = ZoneInfo("America/Santiago")
PHOTO_SAVE_ROOT = os.getenv("PHOTO_SAVE_ROOT", "./photos")
FLASH_SAVE_ROOT = os.getenv("FLASH_SAVE_ROOT", "./Flash_Reportes")
DATA_ROOT = os.getenv("DATA_ROOT", "./data")
os.makedirs(PHOTO_SAVE_ROOT, exist_ok=True)
os.makedirs(FLASH_SAVE_ROOT, exist_ok=True)
os.makedirs(DATA_ROOT, exist_ok=True)

FLASH_GROUP_CHAT_ID = os.getenv("FLASH_GROUP_CHAT_ID", "")
MS_CLIENT_ID = os.getenv("MS_CLIENT_ID", "")
MS_TENANT_ID = os.getenv("MS_TENANT_ID", "common")
MS_SCOPES = ["Files.ReadWrite", "User.Read"]
ONEDRIVE_ROOT = os.getenv("ONEDRIVE_ROOT", "Bot_FotosITO")
TOKEN_CACHE_PATH = os.getenv("TOKEN_CACHE_PATH", "./token_cache.bin")
PORT = int(os.getenv("PORT", "10000"))

PHOTO_PIQUE, PHOTO_FRENTE, PHOTO_MARCO, PHOTO_ETAPA, PHOTO_COMENTARIO = range(1, 6)
PHOTO_PIQUES = ["Túnel Enlace", "Bremen", "Talleres", "Lo Errázuriz", "Román Salinas", "Otros"]
FRENTES_BASE = ["TIE Oriente", "TIE Poniente", "Superficie", "Pique"]
FRENTES_LOE = ["TIE Oriente", "TIE Poniente", "Superficie", "Pique", "TEA", "TEB", "TEC"]
FRENTES_RS = ["Pique", "Superficie"]

FLASH_PIQUE, FLASH_FRENTE, FLASH_EVENTO, FLASH_DETALLE, FLASH_FOTO = range(10, 15)
FLASH_PIQUES = PHOTO_PIQUES
FLASH_FRENTES = ["TIE Oriente", "TIE Poniente", "Galería Oriente", "Galería Poniente", "TEA", "TEB", "TEC", "Tramo B", "Tramo C", "Superficie", "Pique", "Otro"]
EVENTOS_FLASH = ["Desprendimientos", "Inundación", "Sobreexcavación", "Tiempo Frente Abierta", "Falta Alzaprima", "No aplicación de 5cm", "Otro"]

PHOTO_XLSX = os.path.join(DATA_ROOT, "Registro_Fotos.xlsx")
PHOTO_SHEET = "Registro_Fotos"
PHOTO_HEADER = ["ID", "Fecha", "Hora", "FechaHora", "Usuario", "UsuarioID", "Pique", "Frente", "Marco", "Etapa", "Comentario", "Archivo", "RutaOneDrive", "Estado"]
FLASH_XLSX = os.path.join(FLASH_SAVE_ROOT, "Flash_Reportes.xlsx")
FLASH_SHEET = "Registro_Flash"
FLASH_HEADER = ["ID", "Fecha", "Hora", "FechaHora", "Inspector", "UsuarioID", "Pique", "Frente", "Evento", "Detalle", "Foto", "Estado"]

logging.basicConfig(format="%(asctime)s - %(levelname)s - %(message)s", level=logging.INFO)
log = logging.getLogger("BotFotosITO")
PENDING_ONEDRIVE_FLOWS = {}

class Handler(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.end_headers()
        self.wfile.write(b"OK")
    def log_message(self, format, *args):
        return

def start_health():
    server = HTTPServer(("0.0.0.0", PORT), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    log.info(f"Healthcheck OK puerto {PORT}")

def now_chile():
    return datetime.now(CHILE_TZ)

def normalizar_texto(texto: str) -> str:
    texto = str(texto).strip()
    for a, b in {"á":"a","é":"e","í":"i","ó":"o","ú":"u","Á":"A","É":"E","Í":"I","Ó":"O","Ú":"U","ñ":"n","Ñ":"N","ü":"u","Ü":"U"}.items():
        texto = texto.replace(a, b)
    return texto

def clean_filename(text: str) -> str:
    text = normalizar_texto(text)
    text = re.sub(r"[^\w\-]+", "_", text)
    text = re.sub(r"_+", "_", text)
    return text.strip("_")

def get_user_name(user) -> str:
    if not user:
        return "SinUsuario"
    return user.username or user.full_name or str(user.id)

def get_user_label(user) -> str:
    if not user:
        return "Sin usuario"
    if user.username:
        return f"{user.full_name or user.username} (@{user.username})"
    return user.full_name or str(user.id)

def safe_text(value):
    return "" if value is None else str(value)

def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    txt = safe_text(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(txt, fmt).date()
        except Exception:
            pass
    return None

def build_keyboard(items, prefix, cols=2):
    keyboard, row = [], []
    for item in items:
        row.append(InlineKeyboardButton(item, callback_data=f"{prefix}|{item}"))
        if len(row) == cols:
            keyboard.append(row)
            row = []
    if row:
        keyboard.append(row)
    return InlineKeyboardMarkup(keyboard)

def ensure_saved(path: str):
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    if os.path.getsize(path) <= 0:
        raise IOError("Archivo vacío")

def frentes_por_pique(pique: str):
    if pique == "Lo Errázuriz":
        return FRENTES_LOE
    if pique == "Román Salinas":
        return FRENTES_RS
    return FRENTES_BASE

def requiere_marco(frente: str) -> bool:
    return frente in ["TIE Oriente", "TIE Poniente"]

def requiere_etapa(frente: str) -> bool:
    return frente == "Pique"

def make_photo_filename(data: dict):
    extra = ""
    if data.get("marco"):
        extra = f"_MR-{clean_filename(data['marco'])}"
    elif data.get("etapa"):
        extra = f"_ETAPA-{clean_filename(data['etapa'])}"
    return f"{data['fecha_archivo']}_{data['hora_archivo']}_{clean_filename(data['pique'])}_{clean_filename(data['frente'])}{extra}_{clean_filename(data['usuario'])}.jpg"

def make_onedrive_photo_folder(data: dict):
    return f"Fotos/{clean_filename(data['pique'])}/{clean_filename(data['frente'])}"

def aplicar_formato_excel(ws, total_cols):
    header_fill = PatternFill("solid", fgColor="111827")
    header_font = Font(color="FFFFFF", bold=True)
    row_fill = PatternFill("solid", fgColor="F8FAFC")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    max_row = ws.max_row
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=total_cols):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = row_fill
    ws.freeze_panes = "A2"
    max_col_letter = ws.cell(row=1, column=total_cols).column_letter
    ws.auto_filter.ref = f"A1:{max_col_letter}{max_row}"
    widths = {1:18,2:14,3:12,4:22,5:28,6:14,7:20,8:22,9:18,10:18,11:55,12:55,13:45,14:18}
    for col_idx in range(1, total_cols + 1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[letter].width = widths.get(col_idx, 20)
    ws.row_dimensions[1].height = 26
    for r in range(2, max_row + 1):
        ws.row_dimensions[r].height = 36

def crear_excel_si_no_existe(path, sheet_name, headers):
    if os.path.exists(path):
        return
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    aplicar_formato_excel(ws, len(headers))
    wb.save(path)

def guardar_registro_excel(path, sheet_name, headers, row_values, table_name):
    crear_excel_si_no_existe(path, sheet_name, headers)
    wb = load_workbook(path)
    ws = wb[sheet_name]
    ws.append(row_values)
    aplicar_formato_excel(ws, len(headers))
    max_row = ws.max_row
    max_col_letter = ws.cell(row=1, column=len(headers)).column_letter
    table_ref = f"A1:{max_col_letter}{max_row}"
    if max_row >= 2:
        if table_name not in ws.tables:
            tab = Table(displayName=table_name, ref=table_ref)
            style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            tab.tableStyleInfo = style
            ws.add_table(tab)
        else:
            ws.tables[table_name].ref = table_ref
    wb.save(path)

def leer_excel(path, sheet_name):
    if not os.path.exists(path):
        return []
    wb = load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        return []
    headers = [safe_text(x) for x in rows[0]]
    data = []
    for r in rows[1:]:
        item = {}
        for i, h in enumerate(headers):
            item[h] = r[i] if i < len(r) else ""
        data.append(item)
    return data

def guardar_photo_metadata(data: dict):
    guardar_registro_excel(PHOTO_XLSX, PHOTO_SHEET, PHOTO_HEADER, [data["id"], data["fecha"], data["hora"], data["fecha_hora"], data["usuario_label"], data["usuario_id"], data["pique"], data["frente"], data.get("marco", ""), data.get("etapa", ""), data.get("comentario", ""), data["archivo"], data.get("ruta_onedrive", ""), data.get("estado", "Registrado")], "TablaRegistroFotos")

def guardar_flash_xlsx(flash: dict):
    guardar_registro_excel(FLASH_XLSX, FLASH_SHEET, FLASH_HEADER, [flash["id"], flash["fecha"], flash["hora"], flash["fecha_hora"], flash["inspector"], flash["usuario_id"], flash["pique"], flash["frente"], flash["evento"], flash["detalle"], flash["foto"], flash["estado"]], "TablaFlash")

crear_excel_si_no_existe(PHOTO_XLSX, PHOTO_SHEET, PHOTO_HEADER)
crear_excel_si_no_existe(FLASH_XLSX, FLASH_SHEET, FLASH_HEADER)

def load_cache():
    cache = msal.SerializableTokenCache()
    if os.path.exists(TOKEN_CACHE_PATH):
        try:
            with open(TOKEN_CACHE_PATH, "r", encoding="utf-8") as f:
                cache.deserialize(f.read())
        except Exception:
            pass
    return cache

def save_cache(cache):
    if cache.has_state_changed:
        folder = os.path.dirname(TOKEN_CACHE_PATH)
        if folder:
            os.makedirs(folder, exist_ok=True)
        with open(TOKEN_CACHE_PATH, "w", encoding="utf-8") as f:
            f.write(cache.serialize())

def build_msal_app(cache=None):
    return msal.PublicClientApplication(MS_CLIENT_ID, authority=f"https://login.microsoftonline.com/{MS_TENANT_ID}", token_cache=cache)

def get_graph_token():
    if not MS_CLIENT_ID:
        raise RuntimeError("Falta MS_CLIENT_ID en Render.")
    cache = load_cache()
    app = build_msal_app(cache)
    accounts = app.get_accounts()
    if accounts:
        result = app.acquire_token_silent(MS_SCOPES, account=accounts[0])
        if result and "access_token" in result:
            save_cache(cache)
            return result["access_token"]
    raise RuntimeError("Debes ejecutar /onedrive_login")

def upload_to_onedrive(local_path, folder, filename):
    token = get_graph_token()
    root = ONEDRIVE_ROOT.strip("/")
    folder = folder.strip("/")
    remote_path = f"{root}/{folder}/{filename}" if folder else f"{root}/{filename}"
    url = f"https://graph.microsoft.com/v1.0/me/drive/root:/{remote_path}:/content"
    with open(local_path, "rb") as f:
        r = requests.put(url, headers={"Authorization": f"Bearer {token}"}, data=f)
    if r.status_code not in (200, 201):
        raise RuntimeError(f"Graph upload error {r.status_code}: {r.text}")
    return remote_path

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await help_cmd(update, context)

async def menu_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    keyboard = InlineKeyboardMarkup([
        [InlineKeyboardButton("📸 Registrar Foto", callback_data="menu_photo"), InlineKeyboardButton("🚨 Reporte Flash", callback_data="menu_flash")],
        [InlineKeyboardButton("📊 Dashboard", callback_data="menu_dashboard"), InlineKeyboardButton("📈 Estadísticas", callback_data="menu_stats")],
        [InlineKeyboardButton("🔎 Buscar", callback_data="menu_search"), InlineKeyboardButton("👷 Usuarios", callback_data="menu_users")],
        [InlineKeyboardButton("🛰️ Estado", callback_data="menu_status"), InlineKeyboardButton("🧪 Test Grupo", callback_data="menu_testgrupo")],
        [InlineKeyboardButton("📚 Ayuda", callback_data="menu_help")],
    ])
    text = "🚇 *CENTRO DE CONTROL EL6*\n━━━━━━━━━━━━━━━━━━━━━━\n\nSelecciona una función del sistema:"
    target = update.message if update.message else update.callback_query.message
    await target.reply_text(text, reply_markup=keyboard, parse_mode="Markdown")

async def menu_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    if q.data == "menu_photo":
        await q.message.reply_text("📸 *REGISTRO DE FOTO*\n\nEnvía una foto directamente a este chat y el bot iniciará el formulario de metadata.", parse_mode="Markdown")
    elif q.data == "menu_flash":
        await q.message.reply_text("🚨 Para emitir un Reporte Flash usa el comando: /flash")
    elif q.data == "menu_dashboard":
        await dashboard_cmd(update, context)
    elif q.data == "menu_stats":
        await estadisticas_cmd(update, context)
    elif q.data == "menu_search":
        await q.message.reply_text("🔎 *BUSCADOR DE FOTOS*\n\nUsa:\n`/buscar texto`\n\nEjemplos:\n`/buscar MR56`\n`/buscar Bremen`\n`/buscar shotcrete`", parse_mode="Markdown")
    elif q.data == "menu_users":
        await usuarios_cmd(update, context)
    elif q.data == "menu_status":
        await status_cmd(update, context)
    elif q.data == "menu_testgrupo":
        await testgrupo_cmd(update, context)
    elif q.data == "menu_help":
        await help_cmd(update, context)

async def help_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    texto = (
        "🤖 *Inspector Digital EL6*\n━━━━━━━━━━━━━━━━━━━━━━\n\n"
        "Sistema operativo para registrar evidencia en terreno, emitir alertas críticas y respaldar información en OneDrive.\n\n"
        "🧭 *MENÚ PRINCIPAL*\n/menu → abre botones interactivos del sistema\n\n"
        "📸 *REGISTRO DE FOTOS*\nEnvía una foto directamente al bot. El bot pedirá Pique, Frente, N° Marco o Etapa cuando corresponda y Comentario opcional.\n\n"
        "🚨 *INFORME FLASH*\n/flash → emitir evento crítico.\n\n"
        "⚠️ Desprendimientos\n"
        "💧 Inundación\n"
        "⛏️ Sobreexcavación\n"
        "⏱️ Tiempo Frente Abierta\n"
        "🧱 Falta Alzaprima\n"
        "🚧 No aplicación de 5cm\n"
        "➕ Otro\n\n"
        "📊 *CONTROL Y CONSULTAS*\n/dashboard → resumen rápido del día\n/estadisticas → estadísticas últimos 7 días\n/buscar texto → buscar fotos por MR, pique, frente o comentario\n/flashs → últimos reportes Flash\n/usuarios → ranking de usuarios\n\n"
        "🧪 *OPERACIÓN Y SOPORTE*\n/status → estado del bot\n/testgrupo → prueba envío al grupo configurado\n/idchat → obtiene ID del chat o grupo\n/cancel → cancela el flujo actual\n/reset → limpia flujo pendiente\n\n"
        "🔐 *ONEDRIVE*\n/onedrive_login → iniciar autorización\n/onedrive_finish → finalizar autorización\n\n"
        "🛰️ *Recomendación:* Para fotos normales, solo envía la imagen. Para eventos críticos, usa siempre /flash."
    )
    target = update.message if update.message else update.callback_query.message
    await target.reply_text(texto, parse_mode="Markdown")

async def status_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    now = now_chile()
    texto = ("🛰️ *ESTADO DEL SISTEMA*\n━━━━━━━━━━━━━━━━━━━━━━\n\n"
             f"✅ Bot activo\n🕒 Hora Chile: `{now.strftime('%Y-%m-%d %H:%M:%S')}`\n📁 OneDrive Root: `{ONEDRIVE_ROOT}`\n🚨 Grupo Flash ID: `{FLASH_GROUP_CHAT_ID or 'No configurado'}`\n📸 Archivo fotos: `Registro_Fotos.xlsx`\n🚨 Archivo flash: `Flash_Reportes.xlsx`\n🔐 MS_CLIENT_ID: `{'Configurado' if MS_CLIENT_ID else 'No configurado'}`\n🧾 Token cache: `{TOKEN_CACHE_PATH}`")
    target = update.message if update.message else update.callback_query.message
    await target.reply_text(texto, parse_mode="Markdown")

async def testgrupo_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    target = update.message if update.message else update.callback_query.message
    if not FLASH_GROUP_CHAT_ID:
        await target.reply_text("⚠️ FLASH_GROUP_CHAT_ID no está configurado.")
        return
    now = now_chile().strftime("%Y-%m-%d %H:%M:%S")
    try:
        await context.bot.send_message(chat_id=int(FLASH_GROUP_CHAT_ID), text=f"🧪 PRUEBA DE CONEXIÓN\n\n✅ Bot conectado correctamente.\n🕒 Hora Chile: {now}\n🚇 Sistema Reportes Flash EL6 operativo.")
        await target.reply_text("✅ Mensaje de prueba enviado al grupo configurado.")
    except Exception as e:
        await target.reply_text(f"❌ No se pudo enviar al grupo:\n{e}")

async def idchat(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(f"🆔 Chat ID:\n{update.effective_chat.id}\n\n📌 Chat:\n{update.effective_chat.title or 'Chat privado'}")

async def reset_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    target = update.message if update.message else update.callback_query.message
    await target.reply_text("🔄 Flujo limpiado correctamente. Puedes iniciar de nuevo.")

async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    target = update.message if update.message else update.callback_query.message
    await target.reply_text("🛑 Proceso cancelado correctamente.")
    return ConversationHandler.END

async def dashboard_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    target = update.message if update.message else update.callback_query.message
    photos, flashs = leer_excel(PHOTO_XLSX, PHOTO_SHEET), leer_excel(FLASH_XLSX, FLASH_SHEET)
    today = now_chile().date()
    photos_today = [r for r in photos if parse_date(r.get("Fecha")) == today]
    flash_today = [r for r in flashs if parse_date(r.get("Fecha")) == today]
    p_counter = Counter(safe_text(r.get("Pique")) for r in photos_today if safe_text(r.get("Pique")))
    f_counter = Counter(safe_text(r.get("Evento")) for r in flash_today if safe_text(r.get("Evento")))
    texto = f"📊 *DASHBOARD OPERACIONAL EL6*\n━━━━━━━━━━━━━━━━━━━━━━\n\n🕒 Corte: `{now_chile().strftime('%Y-%m-%d %H:%M:%S')}`\n\n📸 Fotos hoy: *{len(photos_today)}*\n🚨 Flash hoy: *{len(flash_today)}*\n📦 Fotos acumuladas: *{len(photos)}*\n🧾 Flash acumulados: *{len(flashs)}*\n\n"
    if p_counter:
        texto += "🏗️ *Fotos hoy por pique:*\n" + "".join(f"• {k}: {v}\n" for k, v in p_counter.most_common(6)) + "\n"
    if f_counter:
        texto += "⚠️ *Flash hoy por evento:*\n" + "".join(f"• {k}: {v}\n" for k, v in f_counter.most_common(6)) + "\n"
    if photos:
        r = photos[-1]
        texto += f"📌 *Última foto:*\n• {safe_text(r.get('FechaHora'))}\n• {safe_text(r.get('Pique'))} / {safe_text(r.get('Frente'))}\n• `{safe_text(r.get('Archivo'))}`\n\n"
    if flashs:
        r = flashs[-1]
        texto += f"🚨 *Último Flash:*\n• {safe_text(r.get('FechaHora'))}\n• {safe_text(r.get('Pique'))} / {safe_text(r.get('Frente'))}\n• {safe_text(r.get('Evento'))}\n"
    await target.reply_text(texto, parse_mode="Markdown")

async def estadisticas_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    target = update.message if update.message else update.callback_query.message
    photos, flashs = leer_excel(PHOTO_XLSX, PHOTO_SHEET), leer_excel(FLASH_XLSX, FLASH_SHEET)
    today = now_chile().date(); start = today - timedelta(days=6)
    photos_7 = [r for r in photos if parse_date(r.get("Fecha")) and start <= parse_date(r.get("Fecha")) <= today]
    flash_7 = [r for r in flashs if parse_date(r.get("Fecha")) and start <= parse_date(r.get("Fecha")) <= today]
    pc = Counter(safe_text(r.get("Pique")) for r in photos_7 if safe_text(r.get("Pique")))
    uc = Counter(safe_text(r.get("Usuario")) for r in photos_7 if safe_text(r.get("Usuario")))
    ec = Counter(safe_text(r.get("Evento")) for r in flash_7 if safe_text(r.get("Evento")))
    texto = f"📈 *ESTADÍSTICAS EL6 — ÚLTIMOS 7 DÍAS*\n━━━━━━━━━━━━━━━━━━━━━━\n\n📅 Periodo: `{start}` a `{today}`\n\n📸 Fotos: *{len(photos_7)}*\n🚨 Flash: *{len(flash_7)}*\n\n🏗️ *Fotos por pique:*\n"
    texto += "".join(f"• {k}: {v}\n" for k, v in pc.most_common(8)) if pc else "• Sin registros\n"
    texto += "\n⚠️ *Flash por evento:*\n" + ("".join(f"• {k}: {v}\n" for k, v in ec.most_common(8)) if ec else "• Sin registros\n")
    texto += "\n👷 *Top usuarios por fotos:*\n" + ("".join(f"• {k}: {v}\n" for k, v in uc.most_common(5)) if uc else "• Sin registros\n")
    await target.reply_text(texto, parse_mode="Markdown")

async def buscar_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = " ".join(context.args).strip()
    if not query:
        await update.message.reply_text("🔎 Usa `/buscar texto`\nEjemplo: `/buscar MR56`", parse_mode="Markdown")
        return
    q = normalizar_texto(query).lower()
    results = []
    for r in leer_excel(PHOTO_XLSX, PHOTO_SHEET):
        haystack = " ".join(safe_text(r.get(k)) for k in ["ID","FechaHora","Usuario","Pique","Frente","Marco","Etapa","Comentario","Archivo"])
        if q in normalizar_texto(haystack).lower():
            results.append(r)
    results = results[-10:]
    if not results:
        await update.message.reply_text(f"🔎 No encontré registros para: `{query}`", parse_mode="Markdown")
        return
    texto = f"🔎 *RESULTADOS*\n━━━━━━━━━━━━━━━━━━━━━━\nConsulta: `{query}`\nMostrando: *{len(results)}*\n\n"
    for i, r in enumerate(reversed(results), 1):
        texto += f"*{i}.* 📸 `{safe_text(r.get('FechaHora'))}`\n🏗️ {safe_text(r.get('Pique'))} / {safe_text(r.get('Frente'))}\n"
        if safe_text(r.get("Marco")): texto += f"🔢 MR: {safe_text(r.get('Marco'))}\n"
        if safe_text(r.get("Etapa")): texto += f"🏷️ Etapa: {safe_text(r.get('Etapa'))}\n"
        texto += f"👷 {safe_text(r.get('Usuario'))}\n📄 `{safe_text(r.get('Archivo'))}`\n\n"
    await update.message.reply_text(texto, parse_mode="Markdown")

async def flashs_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    target = update.message if update.message else update.callback_query.message
    ultimos = leer_excel(FLASH_XLSX, FLASH_SHEET)[-5:]
    if not ultimos:
        await target.reply_text("🚨 Aún no hay reportes Flash registrados.")
        return
    texto = "🚨 *ÚLTIMOS REPORTES FLASH*\n━━━━━━━━━━━━━━━━━━━━━━\n\n"
    for i, r in enumerate(reversed(ultimos), 1):
        texto += f"*{i}.* `{safe_text(r.get('FechaHora'))}`\n🏗️ {safe_text(r.get('Pique'))} / {safe_text(r.get('Frente'))}\n⚠️ {safe_text(r.get('Evento'))}\n👷 {safe_text(r.get('Inspector'))}\n📝 {safe_text(r.get('Detalle'))[:120]}\n\n"
    await target.reply_text(texto, parse_mode="Markdown")

async def usuarios_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    target = update.message if update.message else update.callback_query.message
    photos, flashs = leer_excel(PHOTO_XLSX, PHOTO_SHEET), leer_excel(FLASH_XLSX, FLASH_SHEET)
    photo_users = Counter(safe_text(r.get("Usuario")) for r in photos if safe_text(r.get("Usuario")))
    flash_users = Counter(safe_text(r.get("Inspector")) for r in flashs if safe_text(r.get("Inspector")))
    total = Counter(); total.update(photo_users); total.update(flash_users)
    texto = f"👷 *USUARIOS ACTIVOS — EL6*\n━━━━━━━━━━━━━━━━━━━━━━\n\n📸 Total fotos: *{len(photos)}*\n🚨 Total Flash: *{len(flashs)}*\n\n🏆 *Ranking general:*\n"
    texto += "".join(f"• {u}: {c} registros\n" for u, c in total.most_common(10)) if total else "• Sin registros\n"
    await target.reply_text(texto, parse_mode="Markdown")

async def onedrive_login(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.message:
        return

    if not MS_CLIENT_ID:
        await update.message.reply_text("❌ Falta MS_CLIENT_ID en Render.")
        return

    try:
        cache = load_cache()
        app = build_msal_app(cache)
        flow = app.initiate_device_flow(scopes=MS_SCOPES)
    except Exception as e:
        await update.message.reply_text(f"❌ Error creando flujo OneDrive:\n{e}")
        return

    if "user_code" not in flow:
        await update.message.reply_text(
            f"❌ Error iniciando login de OneDrive:\n{flow.get('error_description', str(flow))}"
        )
        return

    PENDING_ONEDRIVE_FLOWS[str(update.effective_chat.id)] = (app, flow, cache)

    await update.message.reply_text(
        "🔐 AUTORIZACIÓN ONEDRIVE\n\n"
        "1️⃣ Abre el enlace indicado por Microsoft.\n"
        "2️⃣ Ingresa el código.\n"
        "3️⃣ Inicia sesión.\n"
        "4️⃣ Acepta permisos.\n"
        "5️⃣ Cuando Microsoft confirme, vuelve aquí y ejecuta:\n\n"
        "/onedrive_finish\n\n"
        f"{flow['message']}"
    )
async def onedrive_finish(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = str(update.effective_chat.id)
    if chat_id not in PENDING_ONEDRIVE_FLOWS:
        await update.message.reply_text("⚠️ Primero ejecuta /onedrive_login.")
        return
    app, flow, cache = PENDING_ONEDRIVE_FLOWS[chat_id]
    await update.message.reply_text("⏳ Finalizando autorización OneDrive...")
    try:
        result = await asyncio.to_thread(app.acquire_token_by_device_flow, flow)
    except Exception as e:
        await update.message.reply_text(f"❌ Error al finalizar autorización:\n{e}")
        return
    if "access_token" in result:
        save_cache(cache); PENDING_ONEDRIVE_FLOWS.pop(chat_id, None)
        await update.message.reply_text("✅ OneDrive autorizado correctamente.")
    else:
        await update.message.reply_text(f"❌ Error en autorización:\n{result.get('error_description', 'Sin detalle.')}")

async def photo_inicio(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.message or not update.message.photo:
        return ConversationHandler.END
    context.user_data.clear(); file = await update.message.photo[-1].get_file(); now = now_chile(); user = update.message.from_user
    context.user_data["photo"] = {"file":file,"id":now.strftime("%Y%m%d%H%M%S"),"fecha":now.strftime("%Y-%m-%d"),"hora":now.strftime("%H:%M:%S"),"fecha_hora":now.strftime("%Y-%m-%d %H:%M:%S"),"fecha_archivo":now.strftime("%Y%m%d"),"hora_archivo":now.strftime("%H%M%S"),"usuario":get_user_name(user),"usuario_label":get_user_label(user),"usuario_id":str(user.id) if user else "","pique":"","frente":"","marco":"","etapa":"","comentario":"","archivo":"","ruta_onedrive":"","estado":"Registrado"}
    await update.message.reply_text("📸 *FOTO RECIBIDA*\n\nPaso 1 de 4\nSelecciona el *PIQUE*:", reply_markup=build_keyboard(PHOTO_PIQUES, "photo_pique", cols=2), parse_mode="Markdown")
    return PHOTO_PIQUE

async def photo_pique(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer(); data = context.user_data.get("photo")
    if not data: await q.edit_message_text("⚠️ No encuentro foto pendiente. Envía la foto nuevamente."); return ConversationHandler.END
    _, pique = q.data.split("|", 1); data["pique"] = pique
    await q.edit_message_text(f"🏗️ Pique seleccionado: {pique}\n\nPaso 2 de 4\nSelecciona el *FRENTE*:", reply_markup=build_keyboard(frentes_por_pique(pique), "photo_frente", cols=2), parse_mode="Markdown")
    return PHOTO_FRENTE

async def photo_frente(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer(); data = context.user_data.get("photo")
    if not data: await q.edit_message_text("⚠️ No encuentro foto pendiente. Envía la foto nuevamente."); return ConversationHandler.END
    _, frente = q.data.split("|", 1); data["frente"] = frente
    if requiere_marco(frente):
        await q.edit_message_text(f"📍 Frente seleccionado: {frente}\n\nPaso 3 de 4\nIngresa el *N° de Marco*.\nEjemplo: `347`", parse_mode="Markdown")
        return PHOTO_MARCO
    if requiere_etapa(frente):
        await q.edit_message_text(f"📍 Frente seleccionado: {frente}\n\nPaso 3 de 4\nIngresa la *Etapa*.\nEjemplo: `Etapa 3`", parse_mode="Markdown")
        return PHOTO_ETAPA
    return await pedir_comentario(q, context)

async def photo_marco(update: Update, context: ContextTypes.DEFAULT_TYPE):
    data = context.user_data.get("photo")
    if not data: await update.message.reply_text("⚠️ No encuentro foto pendiente. Envía la foto nuevamente."); return ConversationHandler.END
    marco = (update.message.text or "").strip()
    if not marco: await update.message.reply_text("⚠️ Ingresa un N° de Marco válido."); return PHOTO_MARCO
    data["marco"] = marco; return await pedir_comentario(update, context)

async def photo_etapa(update: Update, context: ContextTypes.DEFAULT_TYPE):
    data = context.user_data.get("photo")
    if not data: await update.message.reply_text("⚠️ No encuentro foto pendiente. Envía la foto nuevamente."); return ConversationHandler.END
    etapa = (update.message.text or "").strip()
    if not etapa: await update.message.reply_text("⚠️ Ingresa una etapa válida."); return PHOTO_ETAPA
    data["etapa"] = etapa; return await pedir_comentario(update, context)

async def pedir_comentario(update_or_query, context: ContextTypes.DEFAULT_TYPE):
    keyboard = InlineKeyboardMarkup([[InlineKeyboardButton("Sin comentario", callback_data="photo_sin_comentario")]])
    texto = "📝 Paso 4 de 4\n\nAgrega un comentario opcional.\nEjemplo: `Instalación de marco finalizada sin observaciones.`\n\nSi no deseas agregar comentario, presiona:"
    if hasattr(update_or_query, "edit_message_text"):
        await update_or_query.edit_message_text(texto, reply_markup=keyboard, parse_mode="Markdown")
    else:
        await update_or_query.message.reply_text(texto, reply_markup=keyboard, parse_mode="Markdown")
    return PHOTO_COMENTARIO

async def photo_comentario(update: Update, context: ContextTypes.DEFAULT_TYPE):
    data = context.user_data.get("photo")
    if not data: await update.message.reply_text("⚠️ No encuentro foto pendiente. Envía la foto nuevamente."); return ConversationHandler.END
    data["comentario"] = (update.message.text or "").strip(); return await finalizar_photo(update, context)

async def photo_sin_comentario(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer(); data = context.user_data.get("photo")
    if not data: await q.edit_message_text("⚠️ No encuentro foto pendiente. Envía la foto nuevamente."); return ConversationHandler.END
    data["comentario"] = "Sin comentario"; await q.edit_message_text("📸 Guardando foto y metadata..."); return await finalizar_photo(update, context)

async def finalizar_photo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    data = context.user_data.get("photo")
    if not data:
        return ConversationHandler.END

    nombre = make_photo_filename(data)
    data["archivo"] = nombre

    local_folder = os.path.join(
        PHOTO_SAVE_ROOT,
        clean_filename(data["pique"]),
        clean_filename(data["frente"])
    )

    os.makedirs(local_folder, exist_ok=True)

    local_path = os.path.join(local_folder, nombre)

    try:
        await data["file"].download_to_drive(local_path)
        ensure_saved(local_path)

    except Exception as e:
        context.user_data.clear()

        target = (
            update.message
            if update.message
            else update.callback_query.message
        )

        await target.reply_text(
            f"❌ Error descargando/guardando la foto:\n{e}"
        )

        return ConversationHandler.END

    onedrive_msg = ""
    try:
        data["ruta_onedrive"] = upload_to_onedrive(local_path, make_onedrive_photo_folder(data), nombre); onedrive_msg = "☁️ Foto subida a OneDrive."
    except Exception as e:
        data["ruta_onedrive"] = ""; onedrive_msg = f"⚠️ Foto guardada localmente, pero no se pudo subir a OneDrive:\n{e}"
    try:
        guardar_photo_metadata(data)
        try:
            upload_to_onedrive(PHOTO_XLSX, "Registros", "Registro_Fotos.xlsx"); onedrive_msg += "\n☁️ Registro_Fotos.xlsx actualizado en OneDrive."
        except Exception as e:
            onedrive_msg += f"\n⚠️ Metadata guardada localmente, pero no se pudo subir a OneDrive:\n{e}"
    except Exception as e:
        onedrive_msg += f"\n❌ Error guardando metadata:\n{e}"
    respuesta = f"✅ *FOTO REGISTRADA CORRECTAMENTE*\n━━━━━━━━━━━━━━━━━━━━━━\n\n🆔 ID: `{data['id']}`\n📅 Fecha: `{data['fecha_hora']}`\n👷 Usuario: {data['usuario_label']}\n🏗️ Pique: {data['pique']}\n📍 Frente: {data['frente']}\n"
    if data.get("marco"):
        respuesta += f"🔢 Marco: {data['marco']}\n"

    if data.get("etapa"):
        respuesta += f"🏷️ Etapa: {data['etapa']}\n"

    respuesta += (
        f"📝 Comentario: {data['comentario']}\n\n"
        f"📄 Archivo:\n`{nombre}`\n\n"
        f"{onedrive_msg}"
    )

    target = (
        update.message
        if update.message
        else update.callback_query.message
    )

    await target.reply_text(
        respuesta,
        parse_mode="Markdown"
    )

    context.user_data.clear()

    return ConversationHandler.END
async def flash_inicio(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear(); user = update.message.from_user; now = now_chile()
    context.user_data["flash"] = {"id":now.strftime("%Y%m%d%H%M%S"),"fecha":now.strftime("%Y-%m-%d"),"hora":now.strftime("%H:%M:%S"),"fecha_hora":now.strftime("%Y-%m-%d %H:%M:%S"),"inspector":get_user_label(user),"usuario_id":str(user.id) if user else "","pique":"","frente":"","evento":"","detalle":"","foto":"","foto_path":"","estado":"Emitido"}
    await update.message.reply_text("🚨 *INFORME FLASH EL6*\n\nPaso 1 de 5\nSelecciona el *PIQUE*:", reply_markup=build_keyboard(FLASH_PIQUES, "flash_pique", cols=2), parse_mode="Markdown")
    return FLASH_PIQUE

async def flash_pique(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()

    if "flash" not in context.user_data:
        await q.edit_message_text(
            "⚠️ No encuentro el informe Flash pendiente. Usa /flash nuevamente."
        )
        return ConversationHandler.END

    _, pique = q.data.split("|", 1)
    context.user_data["flash"]["pique"] = pique

    frentes = frentes_por_pique(pique)

    await q.edit_message_text(
        f"🏗️ Pique seleccionado: {pique}\n\n"
        f"Paso 2 de 5\n"
        f"Selecciona el *FRENTE*:",
        reply_markup=build_keyboard(frentes, "flash_frente", cols=2),
        parse_mode="Markdown"
    )

    return FLASH_FRENTE

async def flash_frente(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    if "flash" not in context.user_data: await q.edit_message_text("⚠️ No encuentro el informe Flash pendiente. Usa /flash nuevamente."); return ConversationHandler.END
    _, frente = q.data.split("|", 1); context.user_data["flash"]["frente"] = frente
    await q.edit_message_text(f"📍 Frente seleccionado: {frente}\n\nPaso 3 de 5\nSelecciona el EVENTO:", reply_markup=build_keyboard(EVENTOS_FLASH, "flash_evento", cols=1))
    return FLASH_EVENTO

async def flash_evento(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    if "flash" not in context.user_data: await q.edit_message_text("⚠️ No encuentro el informe Flash pendiente. Usa /flash nuevamente."); return ConversationHandler.END
    _, evento = q.data.split("|", 1); context.user_data["flash"]["evento"] = evento
    await q.edit_message_text(f"⚠️ Evento seleccionado: {evento}\n\nPaso 4 de 5\nEscribe el detalle del informe.\nEjemplo: Frente abierto desde las 10:30 hrs, sin aplicación de shotcrete de 5 cm.")
    return FLASH_DETALLE

async def flash_detalle(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if "flash" not in context.user_data: await update.message.reply_text("⚠️ No encuentro el informe Flash pendiente. Usa /flash nuevamente."); return ConversationHandler.END
    detalle = (update.message.text or "").strip()
    if not detalle: await update.message.reply_text("⚠️ Debes escribir un detalle para el informe Flash."); return FLASH_DETALLE
    context.user_data["flash"]["detalle"] = detalle
    await update.message.reply_text("📸 Paso 5 de 5\n\nEnvía una foto como respaldo.\n\nSi no tienes foto, presiona:", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("Sin foto", callback_data="flash_sin_foto")]]))
    return FLASH_FOTO

async def flash_recibe_foto(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.message or not update.message.photo: return FLASH_FOTO
    if "flash" not in context.user_data: await update.message.reply_text("⚠️ No encuentro el informe Flash pendiente. Usa /flash nuevamente."); return ConversationHandler.END
    flash = context.user_data["flash"]; file = await update.message.photo[-1].get_file()
    nombre = f"FLASH_{clean_filename(flash['fecha_hora'])}_{clean_filename(flash['pique'])}_{clean_filename(flash['frente'])}_{clean_filename(flash['evento'])}.jpg"
    foto_dir = os.path.join(FLASH_SAVE_ROOT, "Fotos"); os.makedirs(foto_dir, exist_ok=True); full = os.path.join(foto_dir, nombre)
    await file.download_to_drive(full); ensure_saved(full); flash["foto"] = nombre; flash["foto_path"] = full
    return await finalizar_flash(update, context)

async def flash_sin_foto(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    if "flash" not in context.user_data: await q.edit_message_text("⚠️ No encuentro el informe Flash pendiente. Usa /flash nuevamente."); return ConversationHandler.END
    context.user_data["flash"]["foto"] = "Sin foto"; context.user_data["flash"]["foto_path"] = ""
    await q.edit_message_text("📋 Generando y enviando Informe Flash...")
    return await finalizar_flash(update, context)

def mensaje_flash(flash: dict) -> str:
    return f"🚨 INFORME FLASH EL6\n\n🆔 ID: {flash['id']}\n📅 Fecha: {flash['fecha_hora']}\n👷 Inspector: {flash['inspector']}\n🏗️ Pique: {flash['pique']}\n📍 Frente: {flash['frente']}\n⚠️ Evento: {flash['evento']}\n\n📝 Detalle:\n{flash['detalle']}\n\n📸 Foto: {flash['foto']}\n\n✅ Registro guardado automáticamente."

async def finalizar_flash(update: Update, context: ContextTypes.DEFAULT_TYPE):
    flash = context.user_data.get("flash")
    if not flash: return ConversationHandler.END
    guardar_flash_xlsx(flash); onedrive_msg = ""
    try:
        upload_to_onedrive(FLASH_XLSX, "Flash_Reportes", "Flash_Reportes.xlsx"); onedrive_msg += "\n☁️ Excel Flash actualizado en OneDrive."
    except Exception as e:
        onedrive_msg += f"\n⚠️ Excel guardado localmente, pero no se pudo subir a OneDrive:\n{e}"
    if flash.get("foto_path"):
        try:
            upload_to_onedrive(flash["foto_path"], "Flash_Reportes/Fotos", flash["foto"]); onedrive_msg += "\n☁️ Foto Flash subida a OneDrive."
        except Exception as e:
            onedrive_msg += f"\n⚠️ Foto guardada localmente, pero no se pudo subir a OneDrive:\n{e}"
    texto = mensaje_flash(flash); enviado_grupo = False
    if FLASH_GROUP_CHAT_ID:
        try:
            if flash.get("foto_path") and os.path.exists(flash["foto_path"]):
                with open(flash["foto_path"], "rb") as img:
                    await context.bot.send_photo(chat_id=int(FLASH_GROUP_CHAT_ID), photo=img, caption=texto)
            else:
                await context.bot.send_message(chat_id=int(FLASH_GROUP_CHAT_ID), text=texto)
            enviado_grupo = True
        except Exception as e:
            log.error(f"Error enviando Flash al grupo: {e}")
    context.user_data.clear(); respuesta = texto + ("\n\n📤 Informe enviado al grupo." if enviado_grupo else "\n\n⚠️ No se pudo enviar al grupo. Revisa FLASH_GROUP_CHAT_ID.") + onedrive_msg
    target = update.message if update.message else update.callback_query.message
    await target.reply_text(respuesta)
    return ConversationHandler.END

def main():
    start_health()
    app = Application.builder().token(BOT_TOKEN).build()
    photo_conv = ConversationHandler(
        entry_points=[MessageHandler(filters.PHOTO, photo_inicio)],
        states={PHOTO_PIQUE:[CallbackQueryHandler(photo_pique, pattern=r"^photo_pique\|")], PHOTO_FRENTE:[CallbackQueryHandler(photo_frente, pattern=r"^photo_frente\|")], PHOTO_MARCO:[MessageHandler(filters.TEXT & ~filters.COMMAND, photo_marco)], PHOTO_ETAPA:[MessageHandler(filters.TEXT & ~filters.COMMAND, photo_etapa)], PHOTO_COMENTARIO:[MessageHandler(filters.TEXT & ~filters.COMMAND, photo_comentario), CallbackQueryHandler(photo_sin_comentario, pattern=r"^photo_sin_comentario$")]},
        fallbacks=[CommandHandler("cancel", cancel), CommandHandler("reset", reset_cmd)], allow_reentry=True)
    flash_conv = ConversationHandler(
        entry_points=[CommandHandler("flash", flash_inicio)],
        states={FLASH_PIQUE:[CallbackQueryHandler(flash_pique, pattern=r"^flash_pique\|")], FLASH_FRENTE:[CallbackQueryHandler(flash_frente, pattern=r"^flash_frente\|")], FLASH_EVENTO:[CallbackQueryHandler(flash_evento, pattern=r"^flash_evento\|")], FLASH_DETALLE:[MessageHandler(filters.TEXT & ~filters.COMMAND, flash_detalle)], FLASH_FOTO:[MessageHandler(filters.PHOTO, flash_recibe_foto), CallbackQueryHandler(flash_sin_foto, pattern=r"^flash_sin_foto$")]},
        fallbacks=[CommandHandler("cancel", cancel), CommandHandler("reset", reset_cmd)], allow_reentry=True)
    for cmd, fn in [("start", start), ("menu", menu_cmd), ("help", help_cmd), ("status", status_cmd), ("testgrupo", testgrupo_cmd), ("idchat", idchat), ("reset", reset_cmd), ("cancel", cancel), ("dashboard", dashboard_cmd), ("estadisticas", estadisticas_cmd), ("buscar", buscar_cmd), ("flashs", flashs_cmd), ("usuarios", usuarios_cmd), ("onedrive_login", onedrive_login), ("onedrivelogin", onedrive_login), ("onedrive_finish", onedrive_finish)]:
        app.add_handler(CommandHandler(cmd, fn))
    app.add_handler(flash_conv)
    app.add_handler(photo_conv)
    app.add_handler(CallbackQueryHandler(menu_callback, pattern=r"^menu_"))
    app.run_polling()

if __name__ == "__main__":
    main()
